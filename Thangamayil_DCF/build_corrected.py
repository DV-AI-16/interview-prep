"""
Build Thangamayil_DCF_corrected.xlsx from the user's original workbook.
Edits are surgical (same layout/format) so the PR diff is readable.
Goal: an honest analyst model -- neither flattering nor punitive.

Fixes:
  1. XNPV: forward CFs only, anchored to transaction date.
  2. Beta: link to own regression (Beta!H10); correct (identity) re-lever; honest
     capital-structure weights (company actual, not peer-avg target).
  3. Share count reconciled to 310.82 lakh (period-end diluted) everywhere.
  4. Terminal value = GGM base; EV/EBITDA exit kept as a labelled cross-check.
  5. Revolver added (new sheet) so projected cash never goes below a min floor;
     balance sheet stays balanced.
  6. (found during audit) Forecast tax computed with tax rate, not interest rate.
"""
import warnings; warnings.filterwarnings("ignore")
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "/projects/sandbox/interview-prep/Thangamayil_DCF/DCF & COMPS of Thangamayil Jwellers.xlsx"
OUT = "/projects/sandbox/interview-prep/Thangamayil_DCF/Thangamayil_DCF_corrected.xlsx"

wb = load_workbook(SRC, data_only=False)
M = wb["Model"]; W = wb["WACC"]; CD = wb["Comps_data"]; SM = wb["Summary"]; AS = wb["Assumption"]

FC = [get_column_letter(c) for c in range(13, 23)]   # M..V  (FY2026..FY2035 forecast)

def note(ws, cell, text):
    ws[cell].comment = Comment("CORRECTED: " + text, "Analyst review")

# ----------------------------------------------------------------- 1. XNPV
M["E159"] = "=XNPV($E$136,$M$152:$X$152,$M$142:$X$142)"
note(M, "E159", "Forward FCFF only (M:X), anchored to transaction date M142. "
     "Previously discounted FY21-25 history and anchored PV to 2021.")

# ----------------------------------------------------------------- 2. Beta / WACC
W["M17"] = "=Beta!H10"                       # regression beta (0.978)
note(W, "M17", "Linked to own 5-yr monthly regression Beta!H10 (~0.98). Was 0.45 hardcoded.")
W["H40"] = "=M17"                            # already levered at current structure
note(W, "H40", "Regression beta is already a levered/equity beta at the current capital "
     "structure, so used directly. Old formula re-levered with (1-equity market return) "
     "instead of (1-tax) -- a variable error.")
W["H24"] = "=K17"                            # honest: company's ACTUAL debt weight (~10%)
note(W, "H24", "Uses the company's ACTUAL capital structure (K17, ~10% debt) rather than the "
     "peer-average 'target' (28%, distorted by Senco/TBZ), which artificially lowered WACC.")

# ----------------------------------------------------------------- 3. Share count
CD["D2"] = 310.82021
note(CD, "D2", "Reconciled to period-end diluted shares (= equity capital 3,108 lakh / FV 10 "
     "= 310.82 lakh), consistent with Model!E137. Was 300.86 (weighted-avg) and inconsistent.")

# ----------------------------------------------------------------- 6. Forecast tax-rate fix
for c in FC:
    M[f"{c}37"] = f"=-{c}36*{c}15"           # EBT x tax rate (row15), not interest rate (row13)
note(M, "M37", "Forecast tax = EBT x TAX RATE (row 15). Was EBT x row 13 (the 7.7% interest "
     "rate), which under-taxed net income. (Does not affect the DCF, which taxes EBIT at E133.)")

# ----------------------------------------------------------------- 5. Revolver sheet
if "Revolver" in wb.sheetnames:
    del wb["Revolver"]
RV = wb.create_sheet("Revolver")
RV.sheet_view.showGridLines = False
HDR = Font(bold=True, color="FFFFFFFF"); FILLH = PatternFill("solid", fgColor="FF1F3864")
BOLD = Font(bold=True); BLUE = Font(color="FF0000FF")
for col, wdt in {"A":34,"B":10,"C":12}.items():
    RV.column_dimensions[col].width = wdt
for c in FC: RV.column_dimensions[c].width = 11
RV["A1"] = "THANGAMAYIL JEWELLERY LIMITED"; RV["A1"].font = BOLD
RV["A2"] = "Revolver / Cash-sweep schedule  (INR in lakh)"; RV["A2"].font = Font(italic=True, color="FF808080")
RV["A4"] = "Minimum cash floor (input)"; RV["A4"].font = BOLD
RV["C4"] = 2000; RV["C4"].font = BLUE
RV["A4"].comment = Comment("Editable assumption: minimum operating cash the business keeps. "
                           "Revolver is drawn to hold cash at or above this level.", "Analyst review")
RV["A5"] = "Revolver interest rate"; RV["A5"].font = BOLD
RV["C5"] = "=Model!M13"
# header row of years
RV["A7"] = "Fiscal year"; RV["A7"].font = HDR; RV["A7"].fill = FILLH
labels = ["Opening revolver balance",
          "Cash available before revolver",
          "Draw / (Repay)",
          "Closing revolver balance",
          "Revolver interest (on opening)"]
rows = {"open":8, "avail":9, "draw":10, "close":11, "int":12}
for r in (8,9,10,11,12):
    RV[f"A{r}"] = labels[r-8]
for c in FC:
    RV[f"{c}7"] = f"=Model!{c}2"; RV[f"{c}7"].font = HDR; RV[f"{c}7"].fill = FILLH
    RV[f"{c}7"].alignment = Alignment(horizontal="center")
prev = None
for c in FC:
    # opening = 0 in first forecast year, else prior closing
    RV[f"{c}8"] = 0 if prev is None else f"={prev}11"
    # cash available before revolver = opening cash + CFO + CFI + (equity + dividends + interest paid)
    RV[f"{c}9"] = (f"=Model!{c}101+Model!{c}82+Model!{c}91"
                   f"+Model!{c}95+Model!{c}96+Model!{c}97")
    # draw if short of floor; repay (down to zero) if surplus
    RV[f"{c}10"] = (f"=IF({c}9<$C$4,$C$4-{c}9,"
                    f"-MIN({c}8,MAX(0,{c}9-$C$4)))")
    RV[f"{c}11"] = f"={c}8+{c}10"
    RV[f"{c}12"] = f"=-{c}8*$C$5"
    prev = c

# ---- wire revolver into Model (forecast cols only) ----
for c in FC:
    M[f"{c}94"] = f"=Revolver!{c}10"                       # financing: draw/(repay)
    M[f"{c}34"] = f"={c}124+Revolver!{c}12"                # IS interest incl revolver (on opening)
    M[f"{c}97"] = f"={c}34"                                # CF interest paid = total interest
    M[f"{c}63"] = f"={c}123+Revolver!{c}11"                # BS debt incl revolver (keeps BS balanced)
note(M, "M94", "Revolver draw/(repay) from the new Revolver sheet (was 0). Sized to keep cash "
     ">= floor. Interest is charged on the opening balance to avoid a circular reference.")
note(M, "M45", "Projected cash now floored by the revolver (was -34,181 in FY26 -- impossible).")
note(M, "M63", "Balance-sheet debt now includes the revolver, so the sheet stays balanced.")

# ----------------------------------------------------------------- 4. Terminal value GGM + cross-check
M["X150"] = "=AA148"                          # GGM as the base terminal
note(M, "X150", "Terminal value = Gordon Growth (AA148). Was AVERAGE(exit-multiple, GGM), which "
     "imported today's rich peer multiple into intrinsic value. Exit multiple kept as cross-check.")
M["Z149"] = "Memo only (not used in base):"
M["AA149"] = "=AVERAGE(AA147:AA148)"          # left as memo
# exit-multiple cross-check value per share
M["D166"] = "Equity Value/Share - EV/EBITDA exit (cross-check)"
M["E166"] = ("=(E163+(AA147-AA148)/(1+E136)^(($X$142-$M$142)/365))/E137")
note(M, "E166", "Cross-check: value per share if terminal used the EV/EBITDA exit multiple "
     "instead of GGM. Drives the exit-multiple sensitivity table.")
# re-point sensitivity table 1 (WACC x exit multiple) to the exit cross-check value
M["I175"] = "=E166"
note(M, "I175", "Sensitivity table 1 (WACC x exit multiple) now reads the EV/EBITDA exit "
     "cross-check (E166); table 2 (WACC x growth) reads the GGM base (E165).")
M["H172"] = "Sensitivities - Exit-multiple cross-check (Rs/Share)"

# ----------------------------------------------------------------- present base case as headline
M["E7"] = "Base Case"
note(M, "E7", "Headline set to Base Case (central estimate). Bear/Bull remain selectable here.")

# make the scenario switch actually work for ALL forecast years (was: only N8 dynamic,
# O8:V8 hardcoded to the Bear path). G7 = 1/2/3 -> Base/Bull/Bear (Scenario rows 5/6/7).
sc_cols = {"N":"I","O":"J","P":"K","Q":"L","R":"M","S":"N","T":"O","U":"P","V":"Q"}
for mc, sc in sc_cols.items():
    M[f"{mc}8"] = f"=CHOOSE($G$7,Scenario!{sc}5,Scenario!{sc}6,Scenario!{sc}7)"
note(M, "N8", "Revenue growth now follows the selected scenario for EVERY forecast year via "
     "CHOOSE(G7,...). Previously only N8 was dynamic; O8:V8 were hardcoded to the Bear path, so "
     "toggling the scenario barely changed the output.")

# ----------------------------------------------------------------- clean template junk
for c in ["C16","D16","E16","F16"]:
    AS[c] = None
note(AS, "C16", "Removed leftover template links (varunbeverages.com).")

# AP day-count 356 -> 365 (historical ratio display)
for c in ["H","I","J","K","L"]:
    M[f"{c}19"] = f"=-{c}$60/{c}$26*365"
note(M, "H19", "Day-count corrected to 365 (was 356).")

# ----------------------------------------------------------------- honest narrative refresh
AS["C34"] = ("WACC ~14.0%. Cost of equity 14.9% from CAPM (Rf 6.68% on the 10-yr G-Sec; ERP "
    "8.42% = equity market return 15.1% - Rf, consistent with Damodaran's India ERP ~8.5%; beta "
    "0.98 from the company's own 5-yr monthly regression vs the Sensex). Post-tax cost of debt "
    "~5.8%. Weights use the company's ACTUAL market-value capital structure (~90% equity / ~10% "
    "debt), not a peer-average target. The earlier 9.4% WACC was understated mainly by a 0.45 "
    "beta and a 28% debt weight, both unsupported.")
AS["C20"] = ("Base Case (DCF intrinsic ~INR 391/share; downside ~89% vs CMP INR 3,497)\n"
    "Revenue growth fades from 20% (FY26 surge) to a 4% terminal rate; margins held near history. "
    "Because the business carries ~140 days of gold inventory, faster growth absorbs heavy working "
    "capital, so incremental ROIC is close to WACC and growth is roughly value-neutral.")
AS["C22"] = ("Bull Case (DCF intrinsic ~INR 353/share)\n"
    "Faster near-term growth (25% fading to 4.5% terminal). Counter-intuitively this is NOT higher "
    "than Base: extra growth ties up more gold inventory at an incremental return near the cost of "
    "capital, so it does not create value. Even the EV/EBITDA exit cross-check (~INR 1,500) stays "
    "far below the market price.")
AS["C24"] = ("Bear Case (DCF intrinsic ~INR 401/share)\n"
    "Slower growth (12% fading to 4%). Value is similar to Base/Bull - the model's value is driven "
    "by the discount rate and the structural working-capital intensity, not by the growth path.")
AS["C32"] = ("On relative value the stock trades at a large premium to peers (EV/Sales ~1.9x vs "
    "~0.9x, EV/EBITDA ~34.7x vs ~16x, P/E ~55x vs ~21x). Applying peer median multiples implies a "
    "fair value of ~INR 1,268 (P/E) to ~INR 1,519 (EV/Sales) per share - still far below the CMP "
    "of INR 3,497. Relative methods inherit peers' already-rich multiples, so they sit above the "
    "intrinsic DCF.")
AS["C36"] = ("All lenses point well below the market price (CMP ~INR 3,497): DCF/GGM intrinsic "
    "~INR 353-401 across Bear/Base/Bull; EV/EBITDA exit-multiple cross-check ~INR 1,200-1,500; "
    "comps ~INR 1,268-1,519. The conclusion is robust to the growth scenario. The stock is priced "
    "for outcomes well beyond what these fundamentals support. (Note: FY26 here is an internal "
    "estimate of ~INR 7,677 Cr built from a projected Q4; the filed FY26 figure is higher "
    "[~INR 8,514 Cr] - refresh from the annual report would lift the base modestly.)")
for cc in ["C20","C22","C24","C32","C34","C36"]:
    AS[cc].alignment = Alignment(wrap_text=True, vertical="top")

# Summary football-field ranges refreshed to corrected outputs (GGM scenario range + comps)
SM["H7"] = 1519; SM["H8"] = 1519; SM["H9"] = 1268; SM["H10"] = 1268      # Comps hi/lo
SM["I7"] = 401;  SM["I8"] = 401;  SM["I9"] = 353;  SM["I10"] = 353        # DCF (GGM) hi/lo
note(SM, "I7", "DCF (GGM) per-share by scenario: Bear 401 / Base 391 / Bull 353 (toggle "
     "Model!E7 to reproduce). Exit-multiple cross-check ~1,214-1,502. Was hardcoded 773/975.")
note(SM, "H7", "Comps implied per-share with reconciled share count: P/E 1,268 ... EV/Sales 1,519.")

# recalc on open (Excel/LibreOffice/Sheets) so data tables & links refresh
try:
    wb.calculation.fullCalcOnLoad = True
except Exception:
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

wb.save(OUT)
print("Saved", OUT)
